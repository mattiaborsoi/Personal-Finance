"""Agent 1: statement parsers, helpers and synthetic fixtures.

No database is needed. Fixtures are generated into ``tmp_path`` (nothing binary is
committed) and every transaction is asserted exactly: inferred years, ledger signs,
foreign spend, per-section ``card_last4`` and statement metadata.
"""

from __future__ import annotations

import csv
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import pytest

from app.services.llm import FakeLLMClient, NullLLMClient
from app.services.parsers.amounts import (
    AmountParts,
    extract_foreign_spend,
    parse_amount,
    parse_amount_parts,
    quantize,
)
from app.services.parsers.base import ParsedTransaction, ParseError, StatementDocument, StatementMetadata
from app.services.parsers.dates import detect_day_first, infer_year, parse_date
from app.services.parsers.llm_extractor import LLMLayoutExtractor, chunk_text, coerce_transaction
from app.services.parsers.metadata import (
    detect_account_type_hint,
    detect_closing_balance,
    detect_closing_date,
    detect_period,
)
from app.services.parsers.pdf_table import PdfTableParser
from app.services.parsers.pdf_text import PdfTextParser, find_card_section, is_header_line
from app.services.parsers.registry import detect_document, parse_statement
from app.services.parsers.tabular import TabularParser
from tests.fixtures import generate

D = Decimal


def txn(day: str, raw: str, amount: str, card: str | None, fx: tuple[str, str] | None = None) -> ParsedTransaction:
    return ParsedTransaction(
        date=date.fromisoformat(day),
        raw_text=raw,
        amount=D(amount),
        card_last4=card,
        foreign_amount=D(fx[1]) if fx else None,
        foreign_currency=fx[0] if fx else None,
    )


AMEX_EXPECTED = [
    txn("2026-07-30", "PAYMENT RECEIVED - THANK YOU", "3384.21", "7715"),
    txn("2026-07-31", "CINEWORLD", "-20.99", "7715"),
    txn("2026-08-04", "ANTHROPIC", "-18.40", "7715", ("USD", "24.00")),
    txn("2026-08-15", "WAITROSE", "-15.81", "7715"),
    txn("2026-08-24", "BRITISH AIRWAYS", "357.99", "7715"),
    txn("2026-07-28", "WAITROSE", "-16.40", "3348"),
    txn("2026-07-30", "ZOOM OCADO", "-37.89", "3348"),
    txn("2026-08-07", "NETFLIX", "-5.99", "3348"),
]

CHECKING_EXPECTED = [
    txn("2026-07-28", "HSBC CARD PYMT", "-3384.21", "4471"),
    txn("2026-08-01", "SALARY ACME", "4500.00", "4471"),
    txn("2026-08-03", "NORTHWIND ENERGY", "-87.27", "4471"),
    txn("2026-08-03", "FIBRELINE BROADBAND", "-30.00", "4471"),
    txn("2026-08-03", "PARTNER TRANSFER CR", "1685.73", "4471"),
    txn("2026-08-05", "ROBINHOOD", "-500.00", "4471"),
]


@pytest.fixture(scope="module")
def fixtures(tmp_path_factory) -> dict[str, Path]:
    return generate.build_all(tmp_path_factory.mktemp("statements"))


# --------------------------------------------------------------------------- #
# Date helpers
# --------------------------------------------------------------------------- #


class TestDates:
    PERIOD = dict(period_start=date(2026, 7, 29), period_end=date(2026, 8, 28))

    @pytest.mark.parametrize(
        "text",
        ["Jul 31", "31 Jul", "31/07/2026", "2026-07-31", "31 Jul 2026", "31 July 2026", "31st July 2026",
         "Jul 31, 2026", "31.07.26", "31-07-2026", "31/07", "2026-07-31 00:00:00"],
    )  # fmt: skip
    def test_formats(self, text):
        assert parse_date(text, **self.PERIOD) == date(2026, 7, 31)

    def test_python_dates_pass_through(self):
        assert parse_date(datetime(2026, 7, 31, 12, 0)) == date(2026, 7, 31)
        assert parse_date(date(2026, 7, 31)) == date(2026, 7, 31)

    def test_not_dates(self):
        assert parse_date("hello") is None
        assert parse_date("") is None
        assert parse_date(None) is None
        assert parse_date("Feb 30", **self.PERIOD) is None
        assert parse_date("Card ending 7715") is None

    def test_year_inference_within_period(self):
        assert parse_date("Aug 04", **self.PERIOD) == date(2026, 8, 4)
        assert parse_date("Jul 28", **self.PERIOD) == date(2026, 7, 28)  # just before the period start

    def test_year_rollover_december_january(self):
        period = dict(period_start=date(2025, 12, 15), period_end=date(2026, 1, 14))
        assert parse_date("Dec 20", **period) == date(2025, 12, 20)
        assert parse_date("Jan 03", **period) == date(2026, 1, 3)
        assert parse_date("20 Dec", **period) == date(2025, 12, 20)
        assert parse_date("20/12", **period) == date(2025, 12, 20)
        assert infer_year(12, period_end=date(2026, 1, 14)) == 2025
        assert infer_year(1, period_end=date(2026, 1, 14)) == 2026

    def test_year_from_start_only(self):
        assert parse_date("Jan 03", period_start=date(2025, 12, 15)) == date(2026, 1, 3)
        assert parse_date("Dec 20", period_start=date(2025, 12, 15)) == date(2025, 12, 20)

    def test_default_year_and_today_fallback(self):
        assert parse_date("Mar 01", default_year=2024) == date(2024, 3, 1)
        today = date.today()
        assert parse_date("Jan 01").year in (today.year, today.year - 1)
        assert parse_date("Jan 01") <= today

    def test_day_first_default_and_fallback(self):
        assert parse_date("03/08/2026") == date(2026, 8, 3)
        assert parse_date("03/08/2026", day_first=False) == date(2026, 3, 8)
        assert parse_date("07/31/2026") == date(2026, 7, 31)  # impossible day-first -> month-first

    def test_detect_day_first(self):
        assert detect_day_first(["01/02/2026", "13/02/2026"]) is True
        assert detect_day_first(["01/02/2026", "02/13/2026"]) is False
        assert detect_day_first(["01/02/2026"]) is True


# --------------------------------------------------------------------------- #
# Amount helpers
# --------------------------------------------------------------------------- #


class TestAmounts:
    @pytest.mark.parametrize(
        "text, expected",
        [
            ("3,384.21", "3384.21"),
            ("-357.99", "-357.99"),
            ("357.99 CR", "357.99"),
            ("357.99CR", "357.99"),
            ("(24.00)", "-24.00"),
            ("£15.81", "15.81"),
            ("-£15.81", "-15.81"),
            ("£-15.81", "-15.81"),
            ("15.81-", "-15.81"),
            ("24.00 DR", "-24.00"),
            ("-357.99 CR", "357.99"),
            ("£ 1,000.00", "1000.00"),
            ("12.5", "12.50"),
            ("1234", "1234.00"),
        ],
    )
    def test_parse_amount(self, text, expected):
        value = parse_amount(text)
        assert isinstance(value, Decimal)
        assert value == D(expected)

    def test_parts(self):
        assert parse_amount_parts("357.99 CR") == AmountParts(D("357.99"), False, "CR")
        assert parse_amount_parts("(24.00)") == AmountParts(D("24.00"), True, None)
        assert parse_amount_parts(3384.21) == AmountParts(D("3384.21"), False, None)
        assert parse_amount_parts(-30) == AmountParts(D("30.00"), True, None)

    def test_not_amounts(self):
        for text in ("abc", "", None, "12.345.6", "CR", "2026-07-31"):
            assert parse_amount(text) is None

    def test_sign_policy(self):
        plain = parse_amount_parts("20.99")
        credit = parse_amount_parts("357.99 CR")
        minus = parse_amount_parts("-357.99")
        assert plain.apply(plain_is_debit=True) == D("-20.99")  # card statement charge
        assert credit.apply(plain_is_debit=True) == D("357.99")
        assert minus.apply(plain_is_debit=True) == D("357.99")  # minus on a card = credit
        assert plain.apply(plain_is_debit=False) == D("20.99")  # signed column
        assert minus.apply(plain_is_debit=False) == D("-357.99")

    def test_quantize_half_up(self):
        assert quantize("2.345") == D("2.35")
        assert quantize("-2.345") == D("-2.35")
        assert quantize("-0.001") == D("0.00")

    @pytest.mark.parametrize(
        "text, cleaned",
        [
            ("ANTHROPIC 18.40 (USD 24.00)", "ANTHROPIC 18.40"),
            ("ANTHROPIC USD 24.00 18.40", "ANTHROPIC 18.40"),
            ("ANTHROPIC 24.00 USD 18.40", "ANTHROPIC 18.40"),
            ("ANTHROPIC 18.40 USD24.00", "ANTHROPIC 18.40"),
        ],
    )
    def test_foreign_spend(self, text, cleaned):
        rest, fx = extract_foreign_spend(text)
        assert rest == cleaned
        assert (fx.currency, fx.amount) == ("USD", D("24.00"))

    def test_foreign_spend_ignores_words_and_base_currency(self):
        assert extract_foreign_spend("ABC LTD 20.00") == ("ABC LTD 20.00", None)
        assert extract_foreign_spend("GBP 20.00 SHOP")[1] is None
        assert extract_foreign_spend("EUR 20.00 SHOP", base_currency="EUR")[1] is None


# --------------------------------------------------------------------------- #
# Metadata detection
# --------------------------------------------------------------------------- #


class TestMetadata:
    @pytest.mark.parametrize(
        "text",
        [
            "Statement period 29 July 2026 to 28 August 2026",
            "Statement period: 29/07/2026 - 28/08/2026",
            "Transactions from 29 July 2026 to 28 August 2026",
            "Period 29 Jul - 28 Aug 2026",
            "from 29 Jul to 28 Aug 2026",
            "29/07/2026 - 28/08/2026",
        ],
    )
    def test_period(self, text):
        assert detect_period(text) == (date(2026, 7, 29), date(2026, 8, 28))

    def test_period_across_year_end_and_absent(self):
        assert detect_period("Statement period 15 Dec 2025 to 14 Jan 2026") == (date(2025, 12, 15), date(2026, 1, 14))
        assert detect_period("no period here 28 Jul - 29 Jul") == (None, None)

    def test_closing_date_and_balance(self):
        assert detect_closing_date("Closing date 28 August 2026") == date(2026, 8, 28)
        assert detect_closing_date("Statement date: 28/08/2026") == date(2026, 8, 28)
        assert detect_closing_date("as at 31 Aug 2026") == date(2026, 8, 31)
        assert detect_closing_date("Fees change as of 1 January 2027. Closing date 28/08/2026") == date(2026, 8, 28)
        assert detect_closing_date("nothing") is None
        assert detect_closing_balance("Closing balance 242.51 CR") == D("242.51")
        assert detect_closing_balance("New balance £1,234.56") == D("1234.56")

    def test_account_type_hint(self):
        card = "Credit Card statement. Minimum payment 25.00. Credit limit 5,000"
        current = "Date Description Paid Out Paid In Balance. Sort code 00-00-00"
        assert detect_account_type_hint(card) == "credit"
        assert detect_account_type_hint(current) == "checking"
        assert detect_account_type_hint("nothing to see") is None

    def test_detect_document_from_filename_for_csv(self, config, tmp_path):
        path = tmp_path / "hsbc_4471_aug.csv"
        path.write_text("Date,Description,Amount\n01/08/2026,SHOP,-1.00\n")
        meta = detect_document(StatementDocument(path), config)
        assert (meta.institution, meta.account_last4, meta.account_type_hint) == ("HSBC", "4471", "checking")

    def test_detect_document_prefers_context_and_config_institution(self, config, tmp_path):
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas

        path = tmp_path / "virgin.pdf"
        c = canvas.Canvas(str(path), pagesize=A4)
        c.drawString(50, 800, "Virgin Money Credit Card")
        c.drawString(50, 780, "Card number **** **** **** 5502")
        c.drawString(50, 760, "Statement date 28/08/2026")
        c.drawString(50, 740, "Balance in 2019 was zero")  # a plain 4-digit group without context
        c.save()
        meta = detect_document(StatementDocument(path), config)
        assert meta.institution == "Virgin"
        assert meta.account_last4 == "5502"
        assert meta.closing_date == date(2026, 8, 28)
        assert meta.account_type_hint == "credit"


# --------------------------------------------------------------------------- #
# Fixture 1: Amex card statement (text PDF)
# --------------------------------------------------------------------------- #


class TestAmexStatement:
    def test_transactions_exact(self, fixtures, config):
        result = parse_statement(fixtures["amex_pdf"], config)
        assert result.parser_name == "pdf_text"
        assert result.transactions == AMEX_EXPECTED
        assert result.warnings == []
        for t in result.transactions:
            assert isinstance(t.amount, Decimal)

    def test_metadata(self, fixtures, config):
        meta = parse_statement(fixtures["amex_pdf"], config).metadata
        assert meta.institution == "Amex"
        assert meta.account_last4 == "7715"
        assert (meta.period_start, meta.period_end) == (date(2026, 7, 29), date(2026, 8, 28))
        assert meta.closing_date == date(2026, 8, 28)
        assert meta.closing_balance == D("242.51")
        assert meta.account_type_hint == "credit"
        assert meta.to_dict()["statement_period"] == "2026-07-29 to 2026-08-28"

    def test_to_dict_matches_blueprint_schema(self, fixtures, config):
        payload = parse_statement(fixtures["amex_pdf"], config).to_dict()
        assert set(payload) == {"statement_metadata", "transactions", "parser", "warnings"}
        first = payload["transactions"][2]
        assert first == {
            "date": "2026-08-04",
            "post_date": None,
            "raw_text": "ANTHROPIC",
            "amount": -18.4,
            "card_last4": "7715",
            "foreign_spend": {"amount": 24.0, "currency": "USD"},
        }

    def test_filename_override_is_used_for_metadata(self, fixtures, config, tmp_path):
        anonymous = tmp_path / "upload.pdf"
        anonymous.write_bytes(fixtures["amex_pdf"].read_bytes())
        result = parse_statement(anonymous, config, filename="statement.pdf")
        assert result.metadata.institution == "Amex"
        assert result.transactions == AMEX_EXPECTED


# --------------------------------------------------------------------------- #
# Fixture 2: HSBC checking statement (table PDF, text PDF, CSV, XLSX)
# --------------------------------------------------------------------------- #


class TestCheckingStatement:
    def _check_meta(self, meta: StatementMetadata) -> None:
        assert meta.institution == "HSBC"
        assert meta.account_last4 == "4471"
        assert (meta.period_start, meta.period_end) == (date(2026, 7, 1), date(2026, 8, 31))
        assert meta.closing_date == date(2026, 8, 31)
        assert meta.account_type_hint == "checking"

    def test_table_pdf(self, fixtures, config):
        result = parse_statement(fixtures["hsbc_table_pdf"], config)
        assert result.parser_name == "pdf_table"
        assert result.transactions == CHECKING_EXPECTED
        self._check_meta(result.metadata)
        assert result.metadata.closing_balance == D("17568.46")

    def test_text_pdf(self, fixtures, config):
        result = parse_statement(fixtures["hsbc_text_pdf"], config)
        assert result.parser_name == "pdf_text"
        assert result.transactions == CHECKING_EXPECTED
        self._check_meta(result.metadata)
        assert result.metadata.closing_balance == D("17568.46")

    def test_table_and_text_variants_parse_identically(self, fixtures, config):
        table = parse_statement(fixtures["hsbc_table_pdf"], config)
        text = parse_statement(fixtures["hsbc_text_pdf"], config)
        assert table.transactions == text.transactions
        assert table.metadata == text.metadata

    def test_each_pdf_parser_alone(self, fixtures, config):
        # Standalone parsers leave card_last4 empty on single-account statements; the
        # registry copies metadata.account_last4 onto every line afterwards.
        doc = StatementDocument(fixtures["hsbc_table_pdf"])
        without_card = [ParsedTransaction(t.date, t.raw_text, t.amount) for t in CHECKING_EXPECTED]
        assert PdfTableParser(config).parse(doc).transactions == without_card
        assert PdfTextParser(config).parse(doc).transactions == without_card
        assert not PdfTableParser(config).can_parse(StatementDocument(fixtures["hsbc_text_pdf"]))

    def test_csv_and_xlsx(self, fixtures, config):
        csv_result = parse_statement(fixtures["checking_csv"], config)
        xlsx_result = parse_statement(fixtures["checking_xlsx"], config)
        assert csv_result.parser_name == xlsx_result.parser_name == "tabular"
        assert csv_result.transactions == CHECKING_EXPECTED
        assert xlsx_result.transactions == CHECKING_EXPECTED
        assert csv_result.metadata == xlsx_result.metadata
        meta = csv_result.metadata
        assert (meta.institution, meta.account_last4, meta.account_type_hint) == ("HSBC", "4471", "checking")
        assert (meta.period_start, meta.period_end) == (date(2026, 7, 28), date(2026, 8, 5))  # from the data
        assert meta.closing_balance == D("17568.46")

    def test_energy_history_csv(self, fixtures, config):
        result = parse_statement(fixtures["energy_history_csv"], config)
        assert result.transactions == [
            txn("2026-05-03", "NORTHWIND ENERGY", "-68.20", "4471"),
            txn("2026-06-03", "NORTHWIND ENERGY", "-68.20", "4471"),
            txn("2026-07-03", "NORTHWIND ENERGY", "-68.20", "4471"),
            txn("2026-08-03", "NORTHWIND ENERGY", "-87.27", "4471"),
        ]

    def test_build_all_names(self, fixtures):
        assert set(fixtures) == {
            "amex_pdf", "hsbc_table_pdf", "hsbc_text_pdf", "checking_csv", "checking_xlsx",
            "energy_history_csv", "prose_pdf",
        }  # fmt: skip
        assert all(p.exists() and p.stat().st_size > 0 for p in fixtures.values())


# --------------------------------------------------------------------------- #
# Tabular parser variants
# --------------------------------------------------------------------------- #


def _write_csv(path: Path, header: list[str], rows: list[list[str]], preamble: list[str] | None = None) -> Path:
    with path.open("w", newline="") as fh:
        for line in preamble or []:
            fh.write(line + "\n")
        writer = csv.writer(fh)
        writer.writerow(header)
        writer.writerows(rows)
    return path


class TestTabular:
    def test_debit_credit_headers_with_preamble(self, config, tmp_path):
        path = _write_csv(
            tmp_path / "export.csv",
            ["Transaction Date", "Narrative", "Debit", "Credit", "Running Balance"],
            [
                ["03/08/2026", "NORTHWIND ENERGY", "87.27", "", "100.00"],
                ["04/08/2026", "REFUND", "", "5.00", "105.00"],
            ],
            preamble=["Account: current", "Exported 01/09/2026", ""],
        )
        result = TabularParser(config).parse(StatementDocument(path))
        assert result.transactions == [
            txn("2026-08-03", "NORTHWIND ENERGY", "-87.27", None),
            txn("2026-08-04", "REFUND", "5.00", None),
        ]

    def test_signed_amount_column(self, config, tmp_path):
        path = _write_csv(
            tmp_path / "monzo.csv",
            ["Date", "Description", "Amount", "Currency"],
            [
                ["2026-08-03", "TESCO", "-12.50", "GBP"],
                ["2026-08-04", "SALARY", "1500.00", "GBP"],
            ],
        )
        result = TabularParser(config).parse(StatementDocument(path))
        assert [t.amount for t in result.transactions] == [D("-12.50"), D("1500.00")]

    def test_amount_column_with_cr_markers_treats_plain_as_debit(self, config, tmp_path):
        path = _write_csv(
            tmp_path / "card.csv",
            ["Date", "Details", "Amount"],
            [["03/08/2026", "TESCO", "12.50"], ["04/08/2026", "REFUND", "12.50 CR"]],
        )
        result = TabularParser(config).parse(StatementDocument(path))
        assert [t.amount for t in result.transactions] == [D("-12.50"), D("12.50")]

    def test_credit_card_export_positive_charges(self, config, tmp_path):
        # An Amex-style export: charges positive, credits negative; the account type comes from the file name.
        path = _write_csv(
            tmp_path / "amex_7715_activity.csv",
            ["Date", "Description", "Card Member", "Amount"],
            [
                ["04/08/2026", "ANTHROPIC", "PRIMARY USER", "18.40"],
                ["24/08/2026", "BRITISH AIRWAYS", "PRIMARY USER", "-357.99"],
            ],
        )
        result = parse_statement(path, config)
        assert result.metadata.account_type_hint == "credit"
        assert [t.amount for t in result.transactions] == [D("-18.40"), D("357.99")]
        assert {t.card_last4 for t in result.transactions} == {"7715"}

    def test_card_export_signed_like_a_bank_account_is_read_as_printed(self, config, tmp_path):
        """HSBC-style card CSV: a minus on purchases, none on the repayment. Nothing may be flipped."""
        path = _write_csv(
            tmp_path / "card.csv",
            ["Date", "Description", "Amount"],
            [
                ["03/08/2026", "TESCO STORES", "-12.50"],
                ["04/08/2026", "PAYMENT RECEIVED - THANK YOU", "300.00"],
                ["05/08/2026", "PRET A MANGER", "-4.20"],
                ["06/08/2026", "REFUND ONLINE SHOP", "9.99"],
            ],
        )
        meta = StatementMetadata(account_type_hint="credit")
        result = TabularParser(config, metadata=meta).parse(StatementDocument(path))
        assert [t.amount for t in result.transactions] == [D("-12.50"), D("300.00"), D("-4.20"), D("9.99")]
        assert any("read as printed" in w for w in result.warnings)

    def test_signed_card_export_without_a_repayment_goes_by_the_majority(self, config, tmp_path):
        path = _write_csv(
            tmp_path / "card.csv",
            ["Date", "Description", "Amount"],
            [
                ["03/08/2026", "TESCO STORES", "-12.50"],
                ["04/08/2026", "CINEMA", "-18.00"],
                ["05/08/2026", "REFUND", "5.00"],
            ],
        )
        meta = StatementMetadata(account_type_hint="credit")
        result = TabularParser(config, metadata=meta).parse(StatementDocument(path))
        assert [t.amount for t in result.transactions] == [D("-12.50"), D("-18.00"), D("5.00")]

    def test_card_style_export_with_a_signed_repayment_is_still_flipped(self, config, tmp_path):
        """Amex-style: charges unsigned, the repayment and refunds with a minus. The usual card reading."""
        path = _write_csv(
            tmp_path / "card.csv",
            ["Date", "Description", "Amount"],
            [
                ["03/08/2026", "TESCO STORES", "12.50"],
                ["04/08/2026", "PAYMENT RECEIVED - THANK YOU", "-300.00"],
                ["05/08/2026", "REFUND", "-5.00"],
                ["06/08/2026", "CINEMA", "18.00"],
            ],
        )
        meta = StatementMetadata(account_type_hint="credit")
        result = TabularParser(config, metadata=meta).parse(StatementDocument(path))
        assert [t.amount for t in result.transactions] == [D("-12.50"), D("300.00"), D("5.00"), D("-18.00")]
        assert not any("read as printed" in w for w in result.warnings)

    def test_amount_plus_type_column(self, config, tmp_path):
        path = _write_csv(
            tmp_path / "typed.csv",
            ["Date", "Type", "Description", "Amount"],
            [["03/08/2026", "DR", "TESCO", "12.50"], ["04/08/2026", "CR", "REFUND", "12.50"]],
        )
        result = TabularParser(config).parse(StatementDocument(path))
        assert [t.amount for t in result.transactions] == [D("-12.50"), D("12.50")]

    def test_money_out_in_month_first_dates_and_foreign_columns(self, config, tmp_path):
        path = _write_csv(
            tmp_path / "us.csv",
            ["Date", "Description", "Money Out", "Money In", "Original Amount", "Original Currency"],
            [
                ["08/13/2026", "ANTHROPIC", "18.40", "", "24.00", "USD"],
                ["08/14/2026", "REFUND", "", "3.00", "", ""],
            ],
        )
        result = TabularParser(config).parse(StatementDocument(path))
        assert result.transactions == [
            txn("2026-08-13", "ANTHROPIC", "-18.40", None, ("USD", "24.00")),
            txn("2026-08-14", "REFUND", "3.00", None),
        ]

    def test_continuation_rows_and_rows_without_amounts(self, config, tmp_path):
        path = _write_csv(
            tmp_path / "cont.csv",
            ["Date", "Description", "Paid Out", "Paid In", "Balance"],
            [
                ["", "Balance brought forward", "", "", "100.00"],
                ["03/08/2026", "SP PIMORONI LTD", "45.90", "", "54.10"],
                ["", "LONDON GB", "", "", ""],
                ["04/08/2026", "NOTHING HERE", "", "", "54.10"],
                ["", "Balance carried forward", "", "", "54.10"],
            ],
        )
        result = TabularParser(config).parse(StatementDocument(path))
        assert result.transactions == [txn("2026-08-03", "SP PIMORONI LTD LONDON GB", "-45.90", None)]
        assert result.metadata.closing_balance == D("54.10")
        assert any("without an amount" in w for w in result.warnings)

    def test_no_header_raises(self, config, tmp_path):
        path = tmp_path / "junk.csv"
        path.write_text("a,b,c\n1,2,3\n")
        with pytest.raises(ParseError):
            TabularParser(config).parse(StatementDocument(path))
        with pytest.raises(ParseError, match="could not extract"):
            parse_statement(path, config)

    def test_xlsx_datetime_and_numeric_cells(self, config, tmp_path):
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.append(["Date", "Description", "Paid Out", "Paid In"])
        ws.append([datetime(2026, 8, 3), "TESCO", 12.5, None])
        ws.append([datetime(2026, 8, 4), "REFUND", None, 3])
        path = tmp_path / "book.xlsx"
        wb.save(str(path))
        result = TabularParser(config).parse(StatementDocument(path))
        assert result.transactions == [
            txn("2026-08-03", "TESCO", "-12.50", None),
            txn("2026-08-04", "REFUND", "3.00", None),
        ]


# --------------------------------------------------------------------------- #
# Text-line parser layouts (plain text, no positions)
# --------------------------------------------------------------------------- #


class TestPdfTextLayouts:
    CARD_META = StatementMetadata(
        period_start=date(2026, 7, 29), period_end=date(2026, 8, 28), account_type_hint="credit"
    )

    def test_card_layout_variants(self, config):
        text = "\n".join(
            [
                "American Express",
                "Card ending 7715",
                "Jul 31  CINEWORLD  20.99",
                "Aug 04  ANTHROPIC  18.40  (USD 24.00)",
                "Aug 05 ANTHROPIC USD 24.00 18.40",
                "Aug 06 Aug 07 SP PIMORONI LTD 45.90",
                "LONDON GB",
                "Aug 24  BRITISH AIRWAYS  357.99 CR",
                "Aug 25  REFUND SHOP  -12.00",
                "Aug 28  PAYMENT RECEIVED - THANK YOU  3,384.21 CR",
                "Total new charges 1,000.00",
                "Supplementary card ending 3348",
                "Jul 28  WAITROSE  16.40",
                "Page 1 of 2",
            ]
        )
        result = PdfTextParser(config).parse_text(text, self.CARD_META)
        assert result.transactions == [
            txn("2026-07-31", "CINEWORLD", "-20.99", "7715"),
            txn("2026-08-04", "ANTHROPIC", "-18.40", "7715", ("USD", "24.00")),
            txn("2026-08-05", "ANTHROPIC", "-18.40", "7715", ("USD", "24.00")),
            ParsedTransaction(
                date(2026, 8, 6),
                "SP PIMORONI LTD LONDON GB",
                D("-45.90"),
                post_date=date(2026, 8, 7),
                card_last4="7715",
            ),
            txn("2026-08-24", "BRITISH AIRWAYS", "357.99", "7715"),
            txn("2026-08-25", "REFUND SHOP", "12.00", "7715"),
            txn("2026-08-28", "PAYMENT RECEIVED - THANK YOU", "3384.21", "7715"),
            txn("2026-07-28", "WAITROSE", "-16.40", "3348"),
        ]

    def test_foreign_spend_on_continuation_line(self, config):
        text = "Aug 04  ANTHROPIC SAN FRANCISCO  18.40\n24.00 USD @ 1.3043\nAug 05 TESCO 5.00"
        result = PdfTextParser(config).parse_text(text, self.CARD_META)
        assert result.transactions[0] == ParsedTransaction(
            date(2026, 8, 4),
            "ANTHROPIC SAN FRANCISCO @ 1.3043",
            D("-18.40"),
            foreign_amount=D("24.00"),
            foreign_currency="USD",
        )
        assert result.transactions[1].raw_text == "TESCO"

    def test_checking_layout_uses_balance_deltas_without_positions(self, config):
        text = "\n".join(
            [
                "HSBC UK",
                "Date Description Paid Out Paid In Balance",
                "Balance brought forward 15,384.21",
                "28 Jul HSBC CARD PYMT 3,384.21 12,000.00",
                "01 Aug SALARY ACME 4,500.00 16,500.00",
                "03 Aug NORTHWIND ENERGY 87.27 16,412.73",
                "03 Aug PARTNER TRANSFER CR 1,685.73 18,098.46",
                "Closing balance 18,098.46",
            ]
        )
        meta = StatementMetadata(
            period_start=date(2026, 7, 1), period_end=date(2026, 8, 31), account_type_hint="checking"
        )
        result = PdfTextParser(config).parse_text(text, meta)
        assert [(t.raw_text, t.amount) for t in result.transactions] == [
            ("HSBC CARD PYMT", D("-3384.21")),
            ("SALARY ACME", D("4500.00")),
            ("NORTHWIND ENERGY", D("-87.27")),
            ("PARTNER TRANSFER CR", D("1685.73")),
        ]
        assert result.metadata.closing_balance == D("18098.46")

    def test_cr_in_description_marks_credit_when_no_other_evidence(self, config):
        text = "03 Aug PARTNER TRANSFER CR 1,685.73\n04 Aug NORTHWIND ENERGY 87.27"
        meta = StatementMetadata(period_end=date(2026, 8, 31), account_type_hint="checking")
        result = PdfTextParser(config).parse_text(text, meta)
        assert [t.amount for t in result.transactions] == [D("1685.73"), D("-87.27")]

    def test_signed_single_amount_column(self, config):
        text = "Date Description Amount Balance\n28 Jul TESCO -12.50 100.00\n29 Jul REFUND 5.00 105.00"
        result = PdfTextParser(config).parse_text(text, StatementMetadata(period_end=date(2026, 8, 31)))
        assert [t.amount for t in result.transactions] == [D("-12.50"), D("5.00")]

    def test_missing_period_warns_and_uses_today(self, config):
        result = PdfTextParser(config).parse_text("Jul 31 CINEWORLD 20.99")
        assert any("statement period not found" in w for w in result.warnings)
        assert result.transactions[0].date <= date.today()

    def test_no_lines_raises(self, config):
        with pytest.raises(ParseError):
            PdfTextParser(config).parse_text("Just some prose.\nNothing else.")

    def test_helpers(self):
        assert is_header_line("Date Description Paid Out Paid In Balance")
        assert is_header_line("Date Description Amount")
        assert not is_header_line("28 Jul HSBC CARD PYMT 3,384.21")
        assert not is_header_line("Your balance in August")
        assert find_card_section("Card ending 7715", set()) == "7715"
        assert find_card_section("Supplementary card ending 3348", set()) == "3348"
        assert find_card_section("Primary card - Primary User 7715", {"7715"}) == "7715"
        assert find_card_section("Card number xxxx-xxxxxx-x7715", set()) == "7715"
        assert find_card_section("Card issued in 2019", set()) is None
        assert find_card_section("HSBC CARD PYMT", set()) is None
        # Virgin Money: the heading is only the cardholder's name and the full card number.
        assert find_card_section("Primary User 0000 00000000 5502", set()) == "5502"
        assert find_card_section("Secondary User 0000 000000006617", set()) == "6617"
        assert find_card_section("A Person 5355 XXXX XXXX 6617", set()) == "6617"
        assert find_card_section("Your new balance 1234 5678", set()) is None


# --------------------------------------------------------------------------- #
# LLM fallback and errors
# --------------------------------------------------------------------------- #

LLM_RESPONSE = {
    "statement_metadata": {
        "institution": "Amex",
        "account_last4": "7715",
        "closing_date": "2026-08-28",
        "statement_period": "2026-07-29 to 2026-08-28",
    },
    "transactions": [
        {
            "date": "2026-08-14",
            "post_date": "2026-08-14",
            "raw_text": "SP PIMORONI LTD LONDON",
            "amount": -45.90,
            "card_last4": "7715",
            "foreign_spend": None,
        },
        {
            "date": "2026-08-04",
            "raw_text": "ANTHROPIC",
            "amount": "-18.40",
            "card_last4": None,
            "foreign_spend": {"amount": 24, "currency": "usd"},
        },
        {"date": "not a date", "raw_text": "BROKEN", "amount": 1},
        "garbage",
    ],
}


class TestLLMFallback:
    def test_prose_pdf_uses_llm_when_available(self, fixtures, config):
        seen: list[tuple[str, str]] = []

        def handler(system: str, user: str) -> dict:
            seen.append((system, user))
            return LLM_RESPONSE

        llm = FakeLLMClient(handler=handler)
        result = parse_statement(fixtures["prose_pdf"], config, llm=llm)
        assert result.parser_name == "llm_layout"
        assert len(llm.calls) == 1
        system, user = seen[0]
        assert "statement_metadata" in system and "negative = money out" in system
        assert "language-model fallback" in user  # the page text was sent
        assert "prose_only.pdf" in user
        pimoroni = ParsedTransaction(
            date(2026, 8, 14), "SP PIMORONI LTD LONDON", D("-45.90"), post_date=date(2026, 8, 14), card_last4="7715"
        )
        assert result.transactions == [pimoroni, txn("2026-08-04", "ANTHROPIC", "-18.40", "7715", ("USD", "24.00"))]
        meta = result.metadata
        assert (meta.institution, meta.account_last4) == ("Amex", "7715")
        assert (meta.period_start, meta.period_end, meta.closing_date) == (
            date(2026, 7, 29), date(2026, 8, 28), date(2026, 8, 28),
        )  # fmt: skip
        assert result.warnings[0].startswith("deterministic parsers found no transactions")
        assert sum("skipped malformed" in w for w in result.warnings) == 2

    def test_llm_not_used_when_deterministic_parsers_succeed(self, fixtures, config):
        llm = FakeLLMClient(fail=True)
        result = parse_statement(fixtures["amex_pdf"], config, llm=llm)
        assert result.parser_name == "pdf_text"
        assert llm.calls == []

    def test_parse_error_without_llm(self, fixtures, config):
        with pytest.raises(ParseError, match="no LLM available"):
            parse_statement(fixtures["prose_pdf"], config)
        with pytest.raises(ParseError, match="no LLM available"):
            parse_statement(fixtures["prose_pdf"], config, llm=NullLLMClient())

    def test_parse_error_when_llm_fails_or_finds_nothing(self, fixtures, config):
        with pytest.raises(ParseError, match="LLM call failed"):
            parse_statement(fixtures["prose_pdf"], config, llm=FakeLLMClient(fail=True))
        empty = FakeLLMClient(responses=[{"statement_metadata": {}, "transactions": []}])
        with pytest.raises(ParseError, match="returned no transactions"):
            parse_statement(fixtures["prose_pdf"], config, llm=empty)

    def test_unsupported_and_missing_files(self, config, tmp_path):
        bad = tmp_path / "notes.txt"
        bad.write_text("hello")
        with pytest.raises(ParseError, match="unsupported file type"):
            parse_statement(bad, config)
        with pytest.raises(ParseError, match="file not found"):
            parse_statement(tmp_path / "missing.pdf", config)

    def test_chunking_and_multi_page_merge(self, config, tmp_path):
        lines = [f"line {i} " + "x" * 50 for i in range(300)]
        chunks = chunk_text("\n".join(lines), max_chars=6000)
        assert len(chunks) > 1
        assert all(len(c) <= 6000 for c in chunks)
        assert "\n".join(chunks).split("\n") == lines
        assert chunk_text("   ") == []

        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas

        path = tmp_path / "two_pages.pdf"
        c = canvas.Canvas(str(path), pagesize=A4)
        c.drawString(50, 800, "Page one prose about the account.")
        c.showPage()
        c.drawString(50, 800, "Page two prose about the account.")
        c.save()
        responses = [
            {
                "statement_metadata": {"institution": "HSBC"},
                "transactions": [{"date": "2026-08-01", "raw_text": "A", "amount": -1}],
            },
            {
                "statement_metadata": {"account_last4": "4471"},
                "transactions": [{"date": "2026-08-02", "raw_text": "B", "amount": 2}],
            },
        ]
        llm = FakeLLMClient(responses=list(responses))
        result = LLMLayoutExtractor(llm, config).parse(StatementDocument(path))
        assert len(llm.calls) == 2
        assert [(t.raw_text, t.amount) for t in result.transactions] == [("A", D("-1.00")), ("B", D("2.00"))]
        assert (result.metadata.institution, result.metadata.account_last4) == ("HSBC", "4471")

    def test_coerce_transaction_infers_years_from_period(self):
        from app.services.parsers.dates import DateContext

        ctx = DateContext(period_start=date(2025, 12, 15), period_end=date(2026, 1, 14))
        item = {"date": "Dec 20", "raw_text": "SHOP", "amount": "£12.34 CR"}
        assert coerce_transaction(item, ctx) == txn("2025-12-20", "SHOP", "12.34", None)
        assert coerce_transaction({"date": "2026-01-03", "raw_text": "", "amount": 1}, ctx) is None
        assert coerce_transaction({"date": "2026-01-03", "raw_text": "X", "amount": None}, ctx) is None


def test_virgin_statement_splits_lines_by_card_section(config, tmp_path):
    path = generate.virgin_statement_pdf(tmp_path / "virgin.pdf")
    result = parse_statement(path, config)
    cards = [(t.card_last4, t.amount) for t in result.transactions]
    assert cards == [
        ("5502", Decimal("250.00")),
        ("5502", Decimal("-42.10")),
        ("5502", Decimal("-6.80")),
        ("6617", Decimal("-8.45")),
        ("6617", Decimal("-12.99")),
    ]
    # The two-digit years are read as years, not left in the description.
    first = result.transactions[1]
    assert first.raw_text == "WAITROSE LONDON"
    assert (first.date, first.post_date) == (date(2026, 8, 15), date(2026, 8, 16))


def test_foreign_spend_written_with_its_rate(config, tmp_path):
    """Virgin prints the foreign amount and rate on the line under a purchase."""
    pdf = generate._Canvas(tmp_path / "card.pdf")
    pdf.line("Virgin Money credit card", size=14, bold=True)
    pdf.line("Statement period: 24/07/2026 - 23/08/2026")
    pdf.cells([(50, "Date"), (105, "Posted"), (160, "Description")], [(450, "Amount")], bold=True)
    pdf.cells([(50, "12 Aug 26"), (105, "14 Aug 26"), (160, "CAFE CENTRAL VIENNA")], [(450, "£10.70")])
    pdf.line("12.50 @ 1.168 ITA", x=160)  # Virgin prints the country, not the currency
    pdf.save()
    (t,) = parse_statement(tmp_path / "card.pdf", config).transactions
    assert t.raw_text == "CAFE CENTRAL VIENNA"
    assert (t.foreign_currency, t.foreign_amount) == ("EUR", Decimal("12.50"))


def test_two_digit_number_after_a_date_stays_in_the_description(config, tmp_path):
    """Only a two-digit number followed by another date is a year; "24 HOUR FITNESS" keeps its 24."""
    pdf = generate._Canvas(tmp_path / "card.pdf")
    pdf.line("Virgin Money credit card", size=14, bold=True)
    pdf.line("Statement period: 24/07/2026 - 23/08/2026")
    pdf.cells([(50, "Date"), (160, "Description")], [(450, "Amount")], bold=True)
    pdf.cells([(50, "16 Aug"), (160, "24 HOUR FITNESS")], [(450, "£20.00")])
    pdf.save()
    result = parse_statement(tmp_path / "card.pdf", config)
    assert [t.raw_text for t in result.transactions] == ["24 HOUR FITNESS"]
